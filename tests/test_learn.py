"""Categories learnt from the ones a person typed into the workbook:
suggested only for what the rules can't place, marked as suggestions, and
never over a category a person typed. Made-up transactions."""
from openpyxl import load_workbook

from statement_tool.categorize import Category
from statement_tool.excel_writer import COL, TRANSACTIONS_SHEET, append_transactions
from statement_tool.learn import CategoryLearner
from statement_tool.models import Transaction

CATEGORIES = [
    Category("Rent received", "income", ["rent"]),
    Category("Insurance", "expense", ["outsurance"], vat=True),
]
# Other transactions in the workbook, which show which words are common.
OTHERS = ["POS PURCHASE SPAR GROCERIES", "SALARY ACME TRADING", "PIZZA PALACE", "ATM WITHDRAWAL"]


def _t(day, desc, debit=None, credit=None):
    return Transaction(client="", bank="Test Bank", statement_period="P1", date=day, description=desc,
                       debit=debit, credit=credit, balance=None, source_file="a.pdf", account="1234567890")


def _column(ws, header):
    return [ws[f"{COL[header]}{key.row}"].value for key in ws[COL["Row Key"]][1:] if key.value]


def test_suggests_the_category_of_clearly_similar_transactions():
    learner = CategoryLearner([
        ("POS PURCHASE ENGEN MEYERSDAL 515715695061", False, "Fuel"),
        ("POS PURCHASE ENGEN MEYERSDAL 516414379794", False, "Fuel"),
        ("POS PURCHASE WOOLWORTHS ROSEBANK 000012210541", False, "Groceries"),
        ("POS PURCHASE SPAR NINA PARK 004630093602", False, "Groceries"),
    ], OTHERS)
    suggestion = learner.suggest("POS PURCHASE ENGEN MEYERSDAL 516209304590", False)
    assert (suggestion.category, suggestion.alike) == ("Fuel", 2)
    # Only common words and numbers (references, card numbers) in common: not alike.
    assert learner.suggest("POS PURCHASE 515715695061", False) is None
    assert learner.suggest("POS PURCHASE CHICKEN LICKEN", False) is None
    # Its most distinctive word (here the branch) was never categorised: left for a person.
    assert learner.suggest("POS PURCHASE ENGEN BRIGHTSTAR 516209304590", False) is None


def test_how_it_was_paid_and_its_month_dont_make_transactions_alike():
    learner = CategoryLearner([
        ("SASOL KAROO 4000*1234 10 MAY | DEBIT CARD PURCHASE FROM", False, "Fuel"),
        ("KXQT | PAYSHAP PAY BY PROXY", False, "Payments to KXQT"),
    ], OTHERS + ["UBER 4000*1234 02 AUG | DEBIT CARD PURCHASE FROM", "PLMRT | PAYSHAP PAY BY PROXY"])
    assert learner.suggest("SASOL KAROO 4000*1234 17 SEP | DEBIT CARD PURCHASE FROM", False).category == "Fuel"
    assert learner.suggest("GAMESTORE 4000*1234 02 MAY | DEBIT CARD PURCHASE FROM", False) is None
    assert learner.suggest("PLMRT | PAYSHAP PAY BY PROXY", False) is None


def test_only_money_going_the_same_way_is_compared():
    learner = CategoryLearner([("TRANSFER FROM J SMITH RENT", True, "Rent received")], OTHERS)
    assert learner.suggest("TRANSFER FROM J SMITH RENT", True).category == "Rent received"
    assert learner.suggest("TRANSFER FROM J SMITH RENT", False) is None


def test_similar_transactions_that_disagree_give_no_suggestion():
    learner = CategoryLearner([("DIGITAL TRANSFER MOM", False, "Family"),
                               ("DIGITAL TRANSFER MOM", False, "Loans")], OTHERS)
    assert learner.suggest("DIGITAL TRANSFER MOM", False) is None


def test_workbook_suggests_and_learns_from_edits_but_never_overrides_them(tmp_path):
    book = tmp_path / "out.xlsx"
    append_transactions(book, [_t("2026-03-01", "POS PURCHASE ENGEN MEYERSDAL 5157", debit=200.0),
                               _t("2026-03-02", "OUTSURANCE PREMIUM", debit=90.0)], CATEGORIES)
    wb = load_workbook(book)
    ws = wb[TRANSACTIONS_SHEET]
    assert _column(ws, "Category Source") == [None, "Rule"]
    ws[f"{COL['Category']}2"] = "Fuel"  # typed by hand
    wb.save(book)

    append_transactions(book, [_t("2026-03-05", "POS PURCHASE ENGEN MEYERSDAL 5160", debit=150.0),
                               _t("2026-03-06", "PIZZA PALACE", debit=80.0)], CATEGORIES)
    wb = load_workbook(book)
    ws = wb[TRANSACTIONS_SHEET]
    assert _column(ws, "Category") == ["Fuel", "Insurance", "Fuel", "Other expenses (uncategorised)"]
    sources = _column(ws, "Category Source")
    assert sources[:2] == ["Set by you", "Rule"]
    assert sources[2] == 'Suggested - like 1 transaction you categorised, e.g. "POS PURCHASE ENGEN MEYERSDAL 5157"'
    assert sources[3] is None
    assert ws.column_dimensions[COL["Suggested Category"]].hidden

    # Typing over a suggestion: kept from then on, and learnt from.
    ws[f"{COL['Category']}4"] = "Fuel & transport"
    wb.save(book)
    append_transactions(book, [], CATEGORIES)
    ws = load_workbook(book)[TRANSACTIONS_SHEET]
    assert _column(ws, "Category")[2] == "Fuel & transport"
    assert _column(ws, "Category Source")[2] == "Set by you"


def test_a_rule_added_later_replaces_a_suggestion_and_vat_follows_the_category(tmp_path):
    book = tmp_path / "out.xlsx"
    append_transactions(book, [
        _t("2026-03-01", "OUTSURANCE CAR", debit=90.0),
        _t("2026-03-02", "CAR INSURANCE KAROO MUTUAL", debit=95.0),
        _t("2026-03-03", "CAR INSURANCE KAROO MUTUAL", debit=95.0),
        _t("2026-03-04", "PIZZA PALACE", debit=80.0),
        _t("2026-03-05", "SPAR GROCERIES", debit=60.0),
        _t("2026-03-06", "SALARY ACME", credit=1000.0),
    ], CATEGORIES)
    wb = load_workbook(book)
    ws = wb[TRANSACTIONS_SHEET]
    # Nothing is learnt from the rule-made "Insurance" category.
    assert _column(ws, "Category")[1:3] == ["Other expenses (uncategorised)"] * 2
    ws[f"{COL['Category']}3"] = "Car cover"  # typed by hand; the next row gets it suggested
    wb.save(book)
    append_transactions(book, [], CATEGORIES)
    ws = load_workbook(book)[TRANSACTIONS_SHEET]
    assert _column(ws, "Category")[:3] == ["Insurance", "Car cover", "Car cover"]
    assert _column(ws, "Category Source")[2].startswith("Suggested")

    rules = [Category("Vehicle insurance", "expense", ["karoo mutual"], vat=True), *CATEGORIES]
    append_transactions(book, [], rules)
    ws = load_workbook(book)[TRANSACTIONS_SHEET]
    # The hand-typed one stays; the suggestion gives way to the new rule.
    assert _column(ws, "Category")[:3] == ["Insurance", "Car cover", "Vehicle insurance"]
    assert _column(ws, "Category Source")[:3] == ["Rule", "Set by you", "Rule"]
    assert _column(ws, "VAT")[2] == "Yes"
