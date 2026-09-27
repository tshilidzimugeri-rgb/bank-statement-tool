"""Workbook writing: dedup across overlapping statements, categories, reports.

Formulas aren't evaluated here (openpyxl can't); these tests pin down the
structure and the stored values. The formulas themselves were checked by
recalculating a real statement's workbook in Excel.
"""
from datetime import date

from openpyxl import load_workbook

from statement_tool.categorize import Category
from statement_tool.excel_writer import (
    CASHFLOW_SHEET,
    COL,
    INCOME_SHEET,
    MONTHLY_SHEET,
    VAT_SUMMARY_SHEET,
    TRANSACTIONS_SHEET,
    append_transactions,
)
from statement_tool.models import Transaction

CATEGORIES = [
    Category("Rent received", "income", ["rent"]),
    Category("Insurance", "expense", ["outsurance"]),
]


def _t(day, desc, debit=None, credit=None, balance=None, source="a.pdf", period="P1"):
    return Transaction(
        client="Dad", bank="Standard Bank", statement_period=period, date=day, description=desc,
        debit=debit, credit=credit, balance=balance, source_file=source, account="10237421516",
    )


def _column(ws, header):
    return [c.value for c in ws[COL[header]][1:] if c.value is not None]


def test_overlapping_statements_do_not_double_count(tmp_path):
    book = tmp_path / "out.xlsx"
    six_month = [
        _t("2026-03-01", "RENT", credit=1000.0, balance=1100.0, source="6m.pdf", period="Mar-Aug"),
        _t("2026-03-05", "OUTSURANCE", debit=200.0, balance=900.0, source="6m.pdf", period="Mar-Aug"),
    ]
    monthly = [  # same transactions, different file and period label
        _t("2026-03-05", "OUTSURANCE", debit=200.0, balance=900.0, source="mar.pdf", period="Mar"),
        _t("2026-04-01", "RENT", credit=1000.0, balance=1900.0, source="apr.pdf", period="Apr"),
    ]
    assert append_transactions(book, six_month, CATEGORIES).added == 2
    result = append_transactions(book, monthly, CATEGORIES)
    assert (result.added, result.skipped_duplicates) == (1, 1)

    ws = load_workbook(book)[TRANSACTIONS_SHEET]
    assert _column(ws, "Description")[:3] == ["RENT", "OUTSURANCE", "RENT"]


def test_rows_sorted_by_date_and_categorised(tmp_path):
    book = tmp_path / "out.xlsx"
    append_transactions(book, [_t("2026-04-02", "OUTSURANCE", debit=5.0, balance=10.0)], CATEGORIES)
    append_transactions(book, [_t("2026-03-01", "MYSTERY", credit=15.0, balance=15.0)], CATEGORIES)

    ws = load_workbook(book)[TRANSACTIONS_SHEET]
    # openpyxl reads date cells back as datetimes.
    assert [d.date() for d in _column(ws, "Date")[:2]] == [date(2026, 3, 1), date(2026, 4, 2)]
    assert _column(ws, "Month")[:2] == ["2026-03", "2026-04"]
    assert _column(ws, "Category")[:2] == ["Other income (uncategorised)", "Insurance"]


def test_hand_edited_category_survives_later_runs_and_reaches_reports(tmp_path):
    book = tmp_path / "out.xlsx"
    append_transactions(book, [_t("2026-03-01", "MYSTERY", debit=50.0, balance=50.0)], CATEGORIES)

    wb = load_workbook(book)
    ws = wb[TRANSACTIONS_SHEET]
    ws[f"{COL['Category']}2"] = "Garden service"
    wb.save(book)

    append_transactions(book, [_t("2026-03-02", "RENT", credit=100.0, balance=150.0)], CATEGORIES)
    wb = load_workbook(book)
    assert _column(wb[TRANSACTIONS_SHEET], "Category")[:2] == ["Garden service", "Rent received"]
    income_labels = [c.value for c in wb[INCOME_SHEET]["A"]]
    assert "Garden service" in income_labels  # unknown category still reported (as an expense)
    expenses_at = income_labels.index("EXPENSES")
    assert income_labels.index("Garden service") > expenses_at


def test_reports_created_in_order_with_statement_balances(tmp_path):
    book = tmp_path / "out.xlsx"
    append_transactions(
        book,
        [
            _t("2026-03-01", "RENT", credit=1000.0, balance=1100.0),
            _t("2026-03-20", "OUTSURANCE", debit=200.0, balance=900.0),
            _t("2026-04-01", "RENT", credit=1000.0, balance=1900.0),
        ],
        CATEGORIES,
    )
    wb = load_workbook(book)
    assert wb.sheetnames == [
        "Review", MONTHLY_SHEET, INCOME_SHEET, CASHFLOW_SHEET, "Category Breakdown", VAT_SUMMARY_SHEET,
        "Mar 2026", "Apr 2026", TRANSACTIONS_SHEET, "Categories",
    ]

    monthly = wb[MONTHLY_SHEET]
    # Month | Opening | In | Out | Net | Closing - opening worked back from
    # the first transaction, then chained from the previous month's close.
    assert [monthly.cell(row=r, column=c).value for r in (4, 5) for c in (1, 2, 6)] == [
        "2026-03", 100.0, 900.0,
        "2026-04", 900.0, 1900.0,
    ]
    assert monthly["C4"].value.startswith("=SUMIFS(Transactions!")

    cash = wb[CASHFLOW_SHEET]
    labels = [c.value for c in cash["A"]]
    assert "Difference (should be 0)" in labels
    stmt_row = labels.index("Closing balance per bank statement") + 1
    assert [cash.cell(row=stmt_row, column=c).value for c in (2, 3, 4)] == [900.0, 1900.0, 1900.0]


def test_new_rule_recategorises_uncategorised_rows_only(tmp_path):
    book = tmp_path / "out.xlsx"
    append_transactions(
        book,
        [_t("2026-03-01", "ROTO FTW", debit=90.0, balance=10.0), _t("2026-03-02", "OUTSURANCE", debit=5.0, balance=5.0)],
        CATEGORIES,
    )
    wb = load_workbook(book)
    wb[TRANSACTIONS_SHEET][f"{COL['Category']}3"] = "Car insurance"  # hand edit
    wb.save(book)

    updated_rules = [Category("Water tanks", "expense", ["roto"]), *CATEGORIES]
    append_transactions(book, [], updated_rules)

    ws = load_workbook(book)[TRANSACTIONS_SHEET]
    assert _column(ws, "Category")[:2] == ["Water tanks", "Car insurance"]


VAT_CATEGORIES = [
    Category("Rent received", "income", ["rent"], vat=True),
    Category("Insurance", "expense", ["outsurance"], vat=True),
    Category("Levies", "expense", ["huurkor"], vat=False),
    Category("Own transfers", "transfer", ["*****"]),
]


def _cells(ws, row, cols="ABCDEFGH"):
    return [ws[f"{c}{row}"].value for c in cols]


def test_month_sheet_separates_income_and_expenses_with_vat(tmp_path):
    book = tmp_path / "out.xlsx"
    append_transactions(
        book,
        [
            _t("2026-07-01", "RENT", credit=1150.0, balance=1150.0),
            _t("2026-07-02", "OUTSURANCE", debit=115.0, balance=1035.0),
            _t("2026-07-03", "HUURKOR", debit=100.0, balance=935.0),
            _t("2026-07-04", "*****0420 TRANSFER", debit=35.0, balance=900.0),  # neither list
        ],
        VAT_CATEGORIES,
    )
    wb = load_workbook(book)
    tx = wb[TRANSACTIONS_SHEET]
    assert _column(tx, "VAT")[:4] == ["Yes", "Yes", "No", "No"]

    ws = wb["Jul 2026"]
    labels = [c.value for c in ws["A"]]
    assert ws["A2"].value == "Income & Expense Report for July 2026 (01/07/26 - 31/07/26)"

    income_hdr = labels.index("INCOME") + 2  # header row is 1-based index + 1
    rent = income_hdr + 1
    assert ws[f"B{rent}"].value == "JUL26R01"
    # Looked up by Row Key (hidden column I -> row number in J), not a fixed row.
    assert ws[f"I{rent}"].value == _column(tx, "Row Key")[0]
    assert ws[f"J{rent}"].value == f"=MATCH($I{rent},{TRANSACTIONS_SHEET}!${COL['Row Key']}:${COL['Row Key']},0)"
    assert ws[f"F{rent}"].value == (
        f"=INDEX({TRANSACTIONS_SHEET}!${COL['Credit']}:${COL['Credit']},$J{rent})"
        f"-INDEX({TRANSACTIONS_SHEET}!${COL['Debit']}:${COL['Debit']},$J{rent})"
    )
    assert ws[f"G{rent}"].value == f'=IF(E{rent}="Yes",F{rent}/1.15,F{rent})'
    assert labels[rent] == "Total income"  # only one income row

    expense_hdr = labels.index("EXPENSES") + 2
    assert [ws[f"B{expense_hdr + i}"].value for i in (1, 2)] == ["JUL26P01", "JUL26P02"]
    assert ws[f"E{expense_hdr + 2}"].value == f"=INDEX({TRANSACTIONS_SHEET}!${COL['VAT']}:${COL['VAT']},$J{expense_hdr + 2})"
    assert ws[f"I{expense_hdr + 2}"].value == _column(tx, "Row Key")[2]
    assert labels[expense_hdr + 2] == "Total expenses"  # the transfer isn't listed
    assert "Difference" in labels

    summary = wb[VAT_SUMMARY_SHEET]
    assert summary["A5"].value == "Jul 2026"
    assert summary["B5"].value.startswith("='Jul 2026'!F")
    assert summary["J5"].value == "=D5-G5"


def test_hand_set_vat_flag_is_kept(tmp_path):
    book = tmp_path / "out.xlsx"
    append_transactions(book, [_t("2026-07-02", "OUTSURANCE", debit=115.0, balance=1.0)], VAT_CATEGORIES)
    wb = load_workbook(book)
    wb[TRANSACTIONS_SHEET][f"{COL['VAT']}2"] = "No"
    wb.save(book)

    append_transactions(book, [_t("2026-07-03", "OUTSURANCE", debit=115.0, balance=2.0)], VAT_CATEGORIES)
    assert _column(load_workbook(book)[TRANSACTIONS_SHEET], "VAT")[:2] == ["No", "Yes"]


def test_review_status_evidence_and_statement_report_are_recorded(tmp_path):
    from statement_tool.excel_writer import REVIEW_FILL, statement_review_row
    from statement_tool.models import StatementResult

    book = tmp_path / "out.xlsx"
    ok = _t("2026-03-01", "RENT", credit=100.0, balance=100.0)
    ok.evidence, ok.confidence = "01/03/2026 RENT 100.00 100.00", 1.0
    unsure = _t("2026-03-02", "MYSTERY", balance=None)
    unsure.unassigned, unsure.status, unsure.confidence = 40.0, "REVIEW_REQUIRED", 0.8
    unsure.check = "the statement doesn't establish whether this money came in or went out"
    result = StatementResult(source_file="a.pdf", ok=True, transactions=[ok, unsure], statement_period="Mar 2026",
                             status="REVIEW_REQUIRED", problems=["something to check"],
                             report={"opening_balance": 0.0, "total_credits": 100.0, "total_debits": 0.0,
                                     "net_movement": 100.0, "computed_closing": 100.0})
    append_transactions(book, [ok, unsure], CATEGORIES, statement=statement_review_row(result))

    wb = load_workbook(book)
    tx = wb[TRANSACTIONS_SHEET]
    assert _column(tx, "Status")[:2] == ["APPROVED", "REVIEW_REQUIRED"]
    assert _column(tx, "Evidence")[0] == "01/03/2026 RENT 100.00 100.00"
    assert _column(tx, "In/Out Not Shown") == [40.0]  # kept, but in neither debit nor credit
    assert tx[f"{COL['Description']}3"].fill.fgColor.rgb.endswith(REVIEW_FILL.fgColor.rgb[-6:])

    review = wb["Review"]
    header = [c.value for c in review[5]]
    line = dict(zip(header, [c.value for c in review[6]]))
    assert (line["Source File"], line["Status"], line["Rows Needing Review"], line["Issues"]) == (
        "a.pdf", "REVIEW_REQUIRED", 1, "something to check")

    # A later statement adds its own line; the first one is kept.
    later = _t("2026-04-01", "RENT", credit=100.0, balance=200.0)
    later.source_file = "b.pdf"
    append_transactions(book, [later], CATEGORIES,
                        statement={"Source File": "b.pdf", "Status": "APPROVED", "Issues": ""})
    review = load_workbook(book)["Review"]
    assert [review.cell(row=r, column=1).value for r in (6, 7)] == ["a.pdf", "b.pdf"]
