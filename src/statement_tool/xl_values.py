"""Stores each formula's value in the workbook next to the formula.

openpyxl writes formulas without their results, so programs that show a
workbook without calculating it (phone and e-mail previews, many viewers)
show those cells empty. This works out every formula the workbook writer
uses - SUM, SUMIFS, COUNTIFS, COUNTIF, COUNTA, IF, AND, OR, ROUND, INDEX,
MATCH and arithmetic - and saves the results in the file, so every viewer
shows the numbers. The formulas stay: Excel still recalculates them when
the workbook is opened, and after any edit.

Checked against Excel itself: tests/test_workbook_formulas.py has Excel
recalculate workbooks and compares every formula cell with the values
stored here.
"""
from __future__ import annotations

import os
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter


class XLError(Exception):
    """An Excel error value (#N/A, #VALUE!, #DIV/0!...)."""

    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class Range:
    sheet: str
    min_row: int
    min_col: int
    max_row: int
    max_col: int


# --- Reading formulas -----------------------------------------------------------------

_SHEET = r"(?:'(?P<qs>(?:[^']|'')+)'|(?P<ps>[A-Za-z_][\w.]*))!"
_CELL = r"\$?[A-Z]{1,3}\$?\d+"
_TOKEN = re.compile(
    r"\s*(?:"
    r"(?P<str>\"(?:[^\"]|\"\")*\")"
    rf"|(?P<ref>(?:{_SHEET})?(?:{_CELL}(?::{_CELL})?|\$?[A-Z]{{1,3}}:\$?[A-Z]{{1,3}}))(?![\w(])"
    r"|(?P<num>\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)"
    r"|(?P<func>[A-Z][A-Z0-9.]*)\("
    r"|(?P<op><>|<=|>=|[-+*/&=<>(),])"
    r")"
)


def _tokens(formula: str) -> list[tuple[str, str]]:
    out, pos = [], 0
    text = formula.strip()
    while pos < len(text):
        m = _TOKEN.match(text, pos)
        if not m or m.end() == pos:
            raise XLError("#NAME?")
        kind = next(k for k in ("str", "ref", "num", "func", "op") if m.group(k) is not None)
        out.append((kind, m.group(kind), m.group("qs") or m.group("ps")))
        pos = m.end()
        while pos < len(text) and text[pos] == " ":
            pos += 1
    return out


def _cell(text: str) -> tuple[int, int]:
    m = re.fullmatch(r"\$?([A-Z]{1,3})\$?(\d+)", text)
    return int(m.group(2)), column_index_from_string(m.group(1))


class _Parser:
    """Excel expression -> nested tuples: ("num", v), ("str", v), ("ref", Range),
    ("func", name, args), ("op", sym, a, b), ("neg", a)."""

    def __init__(self, tokens, sheet: str):
        self.t, self.i, self.sheet = tokens, 0, sheet

    def peek(self, *values):
        if self.i < len(self.t) and self.t[self.i][0] == "op" and self.t[self.i][1] in values:
            return self.t[self.i][1]
        return None

    def take(self):
        self.i += 1
        return self.t[self.i - 1]

    def parse(self):
        node = self.compare()
        if self.i != len(self.t):
            raise XLError("#NAME?")
        return node

    def compare(self):
        node = self.concat()
        while op := self.peek("=", "<>", "<", ">", "<=", ">="):
            self.take()
            node = ("op", op, node, self.concat())
        return node

    def concat(self):
        node = self.add()
        while self.peek("&"):
            self.take()
            node = ("op", "&", node, self.add())
        return node

    def add(self):
        node = self.mul()
        while op := self.peek("+", "-"):
            self.take()
            node = ("op", op, node, self.mul())
        return node

    def mul(self):
        node = self.unary()
        while op := self.peek("*", "/"):
            self.take()
            node = ("op", op, node, self.unary())
        return node

    def unary(self):
        if self.peek("-"):
            self.take()
            return ("neg", self.unary())
        if self.peek("+"):
            self.take()
            return self.unary()
        return self.primary()

    def primary(self):
        kind, text, sheet = self.take()
        if kind == "num":
            return ("num", float(text))
        if kind == "str":
            return ("str", text[1:-1].replace('""', '"'))
        if kind == "ref":
            return ("ref", self.range(text, sheet))
        if kind == "func":
            args = []
            if not self.peek(")"):
                args.append(self.compare())
                while self.peek(","):
                    self.take()
                    args.append(self.compare())
            if not self.peek(")"):
                raise XLError("#NAME?")
            self.take()
            return ("func", text, args)
        if kind == "op" and text == "(":
            node = self.compare()
            if not self.peek(")"):
                raise XLError("#NAME?")
            self.take()
            return node
        raise XLError("#NAME?")

    def range(self, text: str, sheet: str | None) -> Range:
        sheet = sheet.replace("''", "'") if sheet else self.sheet
        body = text.split("!", 1)[1] if "!" in text else text
        if re.fullmatch(r"\$?[A-Z]{1,3}:\$?[A-Z]{1,3}", body):  # whole columns
            a, b = (column_index_from_string(p.strip("$")) for p in body.split(":"))
            return Range(sheet, 1, min(a, b), 0, max(a, b))  # max_row 0: to the sheet's end
        parts = body.split(":")
        (r1, c1), (r2, c2) = _cell(parts[0]), _cell(parts[-1])
        return Range(sheet, min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2))


# --- Working them out ---------------------------------------------------------------------

def _excel_number(value) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, bool):
        return 1.0 if value else 0.0
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, datetime):
        return (value - datetime(1899, 12, 30)).total_seconds() / 86400
    if isinstance(value, date):
        return float((value - date(1899, 12, 30)).days)
    try:
        return float(str(value))
    except ValueError:
        raise XLError("#VALUE!") from None


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _compare(op: str, a, b) -> bool:
    a = "" if a is None and isinstance(b, str) else (0.0 if a is None else a)
    b = "" if b is None and isinstance(a, str) else (0.0 if b is None else b)
    if isinstance(a, (int, float)) and not isinstance(a, bool) and isinstance(b, (int, float)) \
            and not isinstance(b, bool):
        x, y = float(a), float(b)
    elif isinstance(a, str) and isinstance(b, str):
        x, y = a.lower(), b.lower()
    else:
        # Excel orders numbers < text < booleans.
        rank = lambda v: 2 if isinstance(v, bool) else 1 if isinstance(v, str) else 0  # noqa: E731
        x, y = rank(a), rank(b)
    return {"=": x == y, "<>": x != y, "<": x < y, ">": x > y, "<=": x <= y, ">=": x >= y}[op]


def _criterion(criterion):
    """A SUMIFS/COUNTIFS criterion as a test on a cell's value."""
    if not isinstance(criterion, str):
        return lambda v: v is not None and not isinstance(v, str) and _compare("=", v, criterion)
    m = re.match(r"(<>|<=|>=|=|<|>)?(.*)", criterion, re.S)
    op, target = m.group(1) or "=", m.group(2)
    try:
        number = float(target)
    except ValueError:
        number = None
    if number is not None:
        return lambda v: (isinstance(v, (int, float)) and not isinstance(v, bool) and _compare(op, v, number)) or (
            op in ("=", "<>") and isinstance(v, str) and (v.strip() == target.strip()) == (op == "="))
    if target == "":
        return (lambda v: v is None or v == "") if op == "=" else (lambda v: v is not None and v != "")
    if op in ("=", "<>"):
        pattern = _wildcards(target)  # Excel reads * and ? in a text criterion as wildcards
        return lambda v: (isinstance(v, str) and bool(pattern.fullmatch(v))) == (op == "=")
    return lambda v: isinstance(v, str) and _compare(op, v, target)


def _wildcards(target: str) -> re.Pattern:
    """Excel's text matching: case-insensitive, * any text, ? one character,
    ~* ~? ~~ the character itself."""
    out, i = [], 0
    while i < len(target):
        ch = target[i]
        if ch == "~" and i + 1 < len(target) and target[i + 1] in "*?~":
            out.append(re.escape(target[i + 1]))
            i += 2
            continue
        out.append(".*" if ch == "*" else "." if ch == "?" else re.escape(ch))
        i += 1
    return re.compile("".join(out), re.I | re.S)


def _round(value: float, digits: int) -> float:
    exp = Decimal(1).scaleb(-digits)
    return float(Decimal(repr(value)).quantize(exp, rounding=ROUND_HALF_UP))


class _Book:
    def __init__(self, wb):
        self.wb = wb
        self.values: dict[tuple[str, int, int], object] = {}
        self.busy: set = set()

    def cell(self, sheet: str, row: int, col: int):
        key = (sheet, row, col)
        if key in self.values:
            return self.values[key]
        if sheet not in self.wb.sheetnames:
            raise XLError("#REF!")
        raw = self.wb[sheet].cell(row=row, column=col).value
        if isinstance(raw, str) and raw.startswith("="):
            if key in self.busy:
                raise XLError("#REF!")  # circular
            self.busy.add(key)
            try:
                value = self.evaluate(raw, sheet)
            except XLError as err:
                value = err
            finally:
                self.busy.discard(key)
        else:
            value = raw
        self.values[key] = value
        return value

    def evaluate(self, formula: str, sheet: str):
        tree = _Parser(_tokens(formula[1:]), sheet).parse()
        value = self.run(tree)
        if isinstance(value, Range):
            value = self.single(value)
        if isinstance(value, XLError):
            raise value
        return value

    def cells(self, rng: Range) -> list:
        max_row = rng.max_row or self.wb[rng.sheet].max_row
        return [self.cell(rng.sheet, r, c) for r in range(rng.min_row, max_row + 1)
                for c in range(rng.min_col, rng.max_col + 1)]

    def single(self, rng: Range):
        if rng.min_row == rng.max_row and rng.min_col == rng.max_col:
            return self.cell(rng.sheet, rng.min_row, rng.min_col)
        raise XLError("#VALUE!")

    def value(self, node):
        v = self.run(node)
        if isinstance(v, Range):
            v = self.single(v)
        if isinstance(v, XLError):
            raise v
        return v

    def run(self, node):
        kind = node[0]
        if kind in ("num", "str"):
            return node[1]
        if kind == "ref":
            return node[1]
        if kind == "neg":
            return -_excel_number(self.value(node[1]))
        if kind == "op":
            op, a, b = node[1], self.value(node[2]), self.value(node[3])
            if op == "&":
                return _text(a) + _text(b)
            if op in ("=", "<>", "<", ">", "<=", ">="):
                return _compare(op, a, b)
            x, y = _excel_number(a), _excel_number(b)
            if op == "+":
                return x + y
            if op == "-":
                return x - y
            if op == "*":
                return x * y
            if y == 0:
                raise XLError("#DIV/0!")
            return x / y
        return self.function(node[1], node[2])

    def function(self, name: str, args: list):
        if name == "IF":
            test = self.value(args[0])
            chosen = args[1] if _truthy(test) else (args[2] if len(args) > 2 else None)
            return False if chosen is None else self.value(chosen)
        if name in ("AND", "OR"):
            results = [_truthy(self.value(a)) for a in args]
            return all(results) if name == "AND" else any(results)
        if name == "ROUND":
            return _round(_excel_number(self.value(args[0])), int(_excel_number(self.value(args[1]))))
        if name == "SUM":
            total = 0.0
            for a in args:
                v = self.run(a)
                if isinstance(v, Range):
                    for c in self.cells(v):
                        if isinstance(c, XLError):
                            raise c
                        if isinstance(c, (int, float)) and not isinstance(c, bool):
                            total += c
                else:
                    total += _excel_number(self.value(a))
            return total
        if name == "COUNTA":
            return float(sum(1 for a in args for c in self.cells(self.run(a)) if c is not None))
        if name in ("SUMIFS", "COUNTIFS", "COUNTIF"):
            pairs = args[1:] if name == "SUMIFS" else args
            ranges = [self.cells(self.run(pairs[i])) for i in range(0, len(pairs), 2)]
            tests = [_criterion(self.value(pairs[i])) for i in range(1, len(pairs), 2)]
            hits = [all(test(r[k]) for test, r in zip(tests, ranges)) for k in range(len(ranges[0]))]
            if name != "SUMIFS":
                return float(sum(hits))
            values = self.cells(self.run(args[0]))
            return sum(v for v, hit in zip(values, hits)
                       if hit and isinstance(v, (int, float)) and not isinstance(v, bool)) + 0.0
        if name == "MATCH":
            target = self.value(args[0])
            for position, v in enumerate(self.cells(self.run(args[1])), start=1):
                if not isinstance(v, XLError) and v is not None and _compare("=", v, target):
                    return float(position)
            raise XLError("#N/A")
        if name == "INDEX":
            rng = self.run(args[0])
            row = int(_excel_number(self.value(args[1])))
            col = int(_excel_number(self.value(args[2]))) if len(args) > 2 else 1
            max_row = rng.max_row or self.wb[rng.sheet].max_row
            if not (1 <= row <= max_row - rng.min_row + 1 and 1 <= col <= rng.max_col - rng.min_col + 1):
                raise XLError("#REF!")
            v = self.cell(rng.sheet, rng.min_row + row - 1, rng.min_col + col - 1)
            if isinstance(v, XLError):
                raise v
            return 0.0 if v is None else v
        raise XLError("#NAME?")


def _truthy(value) -> bool:
    if isinstance(value, str):
        raise XLError("#VALUE!")
    return bool(_excel_number(value))


# --- Saving the results into the file ----------------------------------------------------

def calculate(path: Path) -> dict[tuple[str, str], object]:
    """{(sheet, "B5"): value} for every formula cell of the workbook."""
    wb = load_workbook(path)
    book = _Book(wb)
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    out[(ws.title, f"{get_column_letter(c.column)}{c.row}")] = book.cell(ws.title, c.row, c.column)
    return out


_FORMULA_CELL = re.compile(r'<c r="(?P<ref>[A-Z]+\d+)"(?P<attrs>[^>]*)><f>(?P<f>.*?)</f>(?:<v>[^<]*</v>|<v\s*/>)?</c>',
                           re.S)


def _cached(value) -> tuple[str, str]:
    """(type attribute, <v> text) for a formula's result."""
    if isinstance(value, XLError):
        return ' t="e"', value.code
    if isinstance(value, bool):
        return ' t="b"', "1" if value else "0"
    if isinstance(value, (int, float)):
        return "", repr(float(value)) if not float(value).is_integer() else str(int(value))
    return ' t="str"', escape(_text(value))


def add_cached_values(path: Path) -> None:
    """Saves each formula's result in the workbook file, beside the formula."""
    values = calculate(path)
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
        infos = {n: z.getinfo(n) for n in z.namelist()}
    workbook_xml = items["xl/workbook.xml"].decode("utf-8")
    rels = items["xl/_rels/workbook.xml.rels"].decode("utf-8")
    targets = {m.group(1): m.group(2) for m in re.finditer(r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels)}
    targets.update({m.group(2): m.group(1) for m in
                    re.finditer(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
    for m in re.finditer(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', workbook_xml):
        name = m.group(1).replace("&amp;", "&").replace("&apos;", "'").replace("&quot;", '"')
        target = targets[m.group(2)].lstrip("/")
        part = target if target.startswith("xl/") else f"xl/{target}"
        xml = items[part].decode("utf-8")

        def fill(cell):
            key = (name, cell.group("ref"))
            if key not in values:
                return cell.group(0)
            attrs = re.sub(r'\s+t="[^"]*"', "", cell.group("attrs"))
            kind, text = _cached(values[key])
            return f'<c r="{cell.group("ref")}"{attrs}{kind}><f>{cell.group("f")}</f><v>{text}</v></c>'

        items[part] = _FORMULA_CELL.sub(fill, xml).encode("utf-8")
    tmp = path.with_name(path.name + ".values")
    try:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
            for name, data in items.items():
                z.writestr(infos[name], data)
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)  # nothing left behind if it couldn't be swapped in
