"""Bank statements as spreadsheets (Excel .xlsx/.xlsm or CSV), read with the
same rules as PDFs.

A spreadsheet can be messy: title and notes above the table, the header
not on the first row, extra columns, totals and notes below, amounts typed
as text ("R1 234.56", "1,234.56 Cr", "(50.00)", "1.234,56"), dates as text
or as Excel dates, debit and credit columns or one signed amount column or
an amount with a separate Cr/Dr column, rows out of date order. The table
is found by its header row; its columns are recognised by their headings.

- The file is read twice, independently: by openpyxl (or the csv module)
  and by reading the .xlsx's XML parts (or the CSV text) directly. A row
  the two readings see differently is marked REVIEW_REQUIRED.
- Every row is checked against the running balance, in the file's order:
  the previous balance plus or minus the amount must give the balance on
  the row. Where the file doesn't say which way money went, the balances
  show it; where they can't, the amount is kept apart, in no total.
- Nothing is guessed or corrected. A row that doesn't add up is kept
  exactly as in the file and marked, with the reason. Columns that aren't
  recognised (a "Balance Check" or page column someone added) are shown in
  the row's evidence but not used.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from ..config import BankLayout, ClientRule, Settings
from ..models import StatementResult
from . import auto_extract
from .amounts import parse_amount
from .bank_detect import detect_bank, extract_statement_period
from .dates import parse_date
from .parser import _find_account_number, _period_range, build_result
from .reading import Row, assess

SPREADSHEET_TYPES = (".xlsx", ".xlsm", ".csv")
HEADER_SEARCH_ROWS = 40

# Column headings (lower case, punctuation and spaces squeezed).
_HEADINGS = {
    "date": ["date", "transaction date", "trans date", "txn date", "posting date", "post date", "value date",
             "date posted", "transaction date time"],
    "description": ["description", "transaction description", "details", "transaction details", "narrative",
                    "particulars", "reference", "memo", "transaction"],
    "amount": ["amount", "transaction amount", "amount r", "amount zar", "value", "amount in r"],
    "debit": ["debit", "debits", "debit amount", "withdrawal", "withdrawals", "money out", "paid out", "dr",
              "amount debited", "payments"],
    "credit": ["credit", "credits", "credit amount", "deposit", "deposits", "money in", "paid in", "cr",
               "amount credited", "receipts"],
    "balance": ["balance", "running balance", "balance r", "balance zar", "account balance", "closing balance"],
    "type": ["type", "dr cr", "cr dr", "debit credit", "credit debit", "dc", "sign", "amount type"],
    "balance_type": ["balance type", "balance dr cr", "balance cr dr", "balance sign"],
}
_DIRECTION_WORDS = {"cr": 1, "credit": 1, "c": 1, "+": 1, "dr": -1, "debit": -1, "d": -1, "-": -1}


def _heading(value) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


@dataclass
class _Table:
    header_row: int
    columns: dict[str, int]  # role -> column index
    headings: list[str]
    rows: list[tuple[int, list]] = field(default_factory=list)  # (row number in the sheet, cells)


# --- Reading the file -----------------------------------------------------------------

def _sheets(path: Path) -> list[tuple[str, list[list]]]:
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        text = next(raw.decode(enc) for enc in ("utf-8-sig", "cp1252", "latin-1") if _decodes(raw, enc))
        return [(path.stem, [list(r) for r in csv.reader(io.StringIO(text), delimiter=_delimiter(text))])]
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        # From row 1 and column A, so row numbers match the second reading's.
        return [(ws.title, [list(r) for r in ws.iter_rows(min_row=1, min_col=1, values_only=True)])
                for ws in wb.worksheets]
    finally:
        wb.close()


# --- The second, independent reading ------------------------------------------------------

_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
       "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
_BUILTIN_DATE_FORMATS = set(range(14, 23)) | set(range(27, 37)) | set(range(45, 48)) | set(range(50, 59))


def _second_sheets(path: Path) -> list[tuple[str, list[list]]]:
    """The same file read without openpyxl or the csv module: the .xlsx's XML
    parts read directly, or the CSV text split by a parser of its own."""
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        text = next(raw.decode(enc) for enc in ("utf-8-sig", "cp1252", "latin-1") if _decodes(raw, enc))
        return [(path.stem, _split_csv(text))]
    import xml.etree.ElementTree as ET
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        shared = []
        if "xl/sharedStrings.xml" in names:
            for si in ET.fromstring(z.read("xl/sharedStrings.xml")).findall("m:si", _NS):
                shared.append("".join(t.text or "" for t in si.iter(f"{{{_NS['m']}}}t")))
        date_styles = _date_styles(ET.fromstring(z.read("xl/styles.xml"))) if "xl/styles.xml" in names else set()
        rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))}
        out = []
        for sheet in ET.fromstring(z.read("xl/workbook.xml")).find("m:sheets", _NS):
            target = rels[sheet.get(f"{{{_NS['r']}}}id")].lstrip("/")
            part = target if target.startswith("xl/") else f"xl/{target}"
            grid: dict[int, dict[int, object]] = {}
            for c in ET.fromstring(z.read(part)).iter(f"{{{_NS['m']}}}c"):
                ref = re.fullmatch(r"([A-Z]+)(\d+)", c.get("r", ""))
                if not ref:
                    continue
                col = 0
                for ch in ref.group(1):
                    col = col * 26 + ord(ch) - 64
                kind, v = c.get("t", "n"), c.find("m:v", _NS)
                if kind == "s" and v is not None:
                    value = shared[int(v.text)]
                elif kind == "inlineStr":
                    value = "".join(t.text or "" for t in c.iter(f"{{{_NS['m']}}}t"))
                elif kind in ("str", "e"):
                    value = v.text if v is not None else None
                elif kind == "b":
                    value = v is not None and v.text == "1"
                elif v is not None and v.text not in (None, ""):
                    number = float(v.text)
                    if int(c.get("s", 0)) in date_styles:
                        value = datetime(1899, 12, 30) + timedelta(days=number)
                    else:
                        value = int(number) if number.is_integer() and "." not in v.text and "E" not in v.text \
                            else number
                else:
                    value = None
                grid.setdefault(int(ref.group(2)), {})[col] = value
            width = max((max(r) for r in grid.values() if r), default=0)
            last = max(grid, default=0)
            out.append((sheet.get("name"), [[grid.get(r, {}).get(k) for k in range(1, width + 1)]
                                            for r in range(1, last + 1)]))
        return out


def _date_styles(styles) -> set[int]:
    custom = {}
    fmts = styles.find("m:numFmts", _NS)
    for f in (fmts if fmts is not None else []):
        code = re.sub(r'"[^"]*"|\[[^\]]*\]|\\.', "", f.get("formatCode", "")).lower()
        custom[int(f.get("numFmtId"))] = bool(re.search(r"[dmy]", code)) and "general" not in code
    xfs = styles.find("m:cellXfs", _NS)
    out = set()
    for i, xf in enumerate(xfs if xfs is not None else []):
        fmt = int(xf.get("numFmtId", 0))
        if fmt in _BUILTIN_DATE_FORMATS or custom.get(fmt):
            out.add(i)
    return out


def _split_csv(text: str) -> list[list]:
    """CSV fields split character by character: quotes may hold the
    delimiter, line breaks and doubled quotes."""
    delimiter = _delimiter(text)
    rows, row, field, quoted, i = [], [], "", False, 0
    while i < len(text):
        ch = text[i]
        if quoted:
            if ch == '"' and text[i + 1:i + 2] == '"':
                field += '"'
                i += 1
            elif ch == '"':
                quoted = False
            else:
                field += ch
        elif ch == '"' and field == "":
            quoted = True
        elif ch == delimiter:
            row.append(field)
            field = ""
        elif ch in "\r\n":
            if ch == "\r" and text[i + 1:i + 2] == "\n":
                i += 1
            row.append(field)
            rows.append(row)
            row, field = [], ""
        else:
            field += ch
        i += 1
    if field or row:
        row.append(field)
        rows.append(row)
    return rows


def _delimiter(text: str) -> str:
    """The character separating a CSV's columns: the one on the most of its
    first lines, then the most often. (A file with decimal commas,
    "1 234,56;Cr", has commas in its amounts too, but fewer.)"""
    lines = [ln for ln in text.splitlines()[:40] if ln.strip()]
    return max(",;\t|", key=lambda d: (sum(d in ln for ln in lines), sum(ln.count(d) for ln in lines)))


def _decodes(raw: bytes, encoding: str) -> bool:
    try:
        raw.decode(encoding)
        return True
    except UnicodeDecodeError:
        return False


def _text_of(cells: list[list]) -> str:
    return "\n".join(" ".join(_cell_text(v) for v in row if v not in (None, "")) for row in cells)


def _cell_text(value) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d") if not (value.hour or value.minute) else value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return repr(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def _find_table(cells: list[list]) -> _Table | None:
    """The header row: a date column and money columns, within the first rows."""
    for i, row in enumerate(cells[:HEADER_SEARCH_ROWS]):
        headings = [_heading(v) for v in row]
        columns: dict[str, int] = {}
        for role, names in _HEADINGS.items():
            for j, h in enumerate(headings):
                if h in names and j not in columns.values() and role not in columns:
                    columns[role] = j
        money = "amount" in columns or "debit" in columns or "credit" in columns
        if "date" in columns and money and ("balance" in columns or "description" in columns):
            table = _Table(i, columns, [_cell_text(v) for v in row])
            for n, r in enumerate(cells[i + 1:], start=i + 2):
                if any(v not in (None, "") for v in r):
                    table.rows.append((n, list(r) + [None] * (len(row) - len(r))))
            _check_direction_columns(table)
            return table
    return None


def _check_direction_columns(table: _Table) -> None:
    """A "Type" column is a Cr/Dr column only if that's all it holds; else it
    describes the transaction and joins the description."""
    for role in ("type", "balance_type"):
        if role not in table.columns:
            continue
        values = {_heading(r[table.columns[role]]) for _, r in table.rows} - {""}
        if not values <= set(_DIRECTION_WORDS):
            table.columns[f"{role}_text"] = table.columns.pop(role)


# --- One row --------------------------------------------------------------------------

@dataclass
class _Cells:
    """A row's values, straight from its columns."""
    date: str | None  # ISO, when the cell is a date or reads as one
    date_text: str
    description: str
    amount: float | None  # size
    direction: int | None  # +1 in, -1 out, None when the row doesn't say
    balance: float | None
    raw_amount: str  # as in the file, when it isn't a number
    raw_balance: str
    evidence: str
    note: str = ""  # a disagreement between the two readings of the file


def _money(value) -> tuple[float | None, int | None, str]:
    """(value, sign from the cell itself, text when unreadable)."""
    if value in (None, ""):
        return None, None, ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value), (-1 if value < 0 else None), ""
    text = str(value).strip()
    parsed = parse_amount(text)
    if parsed is None:
        return None, None, text
    lower = text.lower()
    sign = -1 if parsed < 0 else (1 if re.search(r"cr\s*$", lower) else None)
    return parsed, sign, ""


def _row(table: _Table, cells: list) -> _Cells:
    col = table.columns

    def at(role):
        return cells[col[role]] if role in col and col[role] < len(cells) else None

    raw_date = at("date")
    if isinstance(raw_date, datetime):
        iso, date_text = raw_date.date().isoformat(), raw_date.date().isoformat()
    elif isinstance(raw_date, date):
        iso, date_text = raw_date.isoformat(), raw_date.isoformat()
    else:
        date_text = _cell_text(raw_date) if raw_date not in (None, "") else ""
        iso = parse_date(date_text) if date_text else None
    words = [_cell_text(at(r)) for r in ("description", "type_text") if at(r) not in (None, "")]
    description = " | ".join(w for w in words if w)

    amount = direction = None
    raw_amount = ""
    debit, debit_sign, debit_raw = _money(at("debit"))
    credit, credit_sign, credit_raw = _money(at("credit"))
    if debit not in (None, 0.0):
        amount, direction = abs(debit), -1
    elif credit not in (None, 0.0):
        amount, direction = abs(credit), 1
    elif "amount" in col:
        value, sign, raw_amount = _money(at("amount"))
        if value is not None:
            amount, direction = abs(value), sign
    else:
        amount = 0.0 if (debit == 0.0 or credit == 0.0) else None
    raw_amount = raw_amount or debit_raw or credit_raw
    stated = _DIRECTION_WORDS.get(_heading(at("type"))) if "type" in col else None
    if stated is not None:
        direction = stated

    balance, balance_sign, raw_balance = _money(at("balance"))
    stated_balance = _DIRECTION_WORDS.get(_heading(at("balance_type"))) if "balance_type" in col else None
    if balance is not None and stated_balance == -1:
        balance = -abs(balance)

    evidence = " | ".join(f"{table.headings[j] or '?'}: {_cell_text(v)}" for j, v in enumerate(cells)
                          if v not in (None, "") and j < len(table.headings))
    return _Cells(iso, date_text, description, amount, direction, balance, raw_amount, raw_balance, evidence)


_OPENING = re.compile(r"opening balance|balance brought forward|brought forward|balance b/?f\b|previous balance",
                      re.IGNORECASE)
MAX_UNKNOWN = 10  # rows without a balance whose direction is worked out together


def _chain(source: list[_Cells], opening: float | None) -> list[Row]:
    """Each row checked against the running balance, in the file's order.
    A row's amount is confirmed when the previous balance plus or minus it
    gives the balance printed on it (rows printed without a balance are
    confirmed together by the next balance). Where the file doesn't say
    which way money went, the balances show it; where they can't, the
    amount is kept apart, in no total. A row that doesn't add up is kept as
    in the file and marked, and checking carries on from its printed
    balance, so one bad row doesn't spill onto the next."""
    rows: list[Row] = []
    previous = opening
    waiting: list[tuple[_Cells, Row]] = []  # amounts printed without a balance

    def new_row(c: _Cells, direction: int | None, check: str = "") -> Row:
        amount = c.amount
        check = "; ".join(x for x in (check, c.note) if x)
        return Row(c.date or (c.date_text or None), c.description or "(no description)",
                   amount if direction == -1 else None, amount if direction == 1 else None, c.balance,
                   amount if direction is None else None, c.evidence, check)

    def settle(balance_row: _Cells | None) -> None:
        """Rows waiting for a balance, confirmed (or not) by balance_row's."""
        nonlocal previous, waiting
        group = waiting + ([(balance_row, None)] if balance_row is not None else [])
        waiting = []
        if not group:
            return
        end = balance_row.balance if balance_row is not None else None
        unknown = [i for i, (c, _) in enumerate(group) if c.direction is None and c.amount]
        fits = []
        if previous is not None and end is not None and len(unknown) <= MAX_UNKNOWN:
            for combo in range(2 ** len(unknown)):
                signs = {i: (1 if combo >> k & 1 else -1) for k, i in enumerate(unknown)}
                total = sum((signs.get(i, c.direction) or 0) * (c.amount or 0) for i, (c, _) in enumerate(group))
                if abs(previous + total - end) <= 0.005:
                    fits.append(signs)
        for i, (c, _) in enumerate(group):
            if len(fits) == 1:
                rows.append(new_row(c, fits[0].get(i, c.direction)))
            elif previous is None:
                rows.append(new_row(c, c.direction, "no earlier balance in the file to confirm this amount"))
            elif end is None:
                rows.append(new_row(c, c.direction, "no balance after it in the file to confirm this amount"))
            elif len(fits) > 1:
                rows.append(new_row(c, c.direction, "more than one reading of the rows before this balance "
                                                    "adds up - check them"))
            else:
                expected = previous + sum((c2.direction or 0) * (c2.amount or 0) for c2, _ in group)
                rows.append(new_row(c, c.direction,
                                    "does not follow from the previous balance - a row may be missing or misread"
                                    + (f" (the amounts give {expected:,.2f}, the file says {end:,.2f})"
                                       if all(c2.direction is not None for c2, _ in group) else "")))
        if end is not None:
            previous = end

    for c in source:
        if not c.amount and not c.raw_amount:
            if c.balance is not None and (_OPENING.search(c.description) or previous is None) and not rows:
                previous = c.balance  # an opening balance line: where the running balance starts
            continue  # a note or total line moves no money
        if c.date is None and c.balance is None and not c.raw_balance and not c.date_text:
            continue  # an undated total below the table ("Count of rows: 836")
        if c.amount is None:  # an amount that isn't a number
            settle(None) if waiting else None
            rows.append(new_row(c, None, f"the amount isn't a readable number ({c.raw_amount})"))
            previous = c.balance if c.balance is not None else None
            continue
        if c.balance is None and not c.raw_balance:
            waiting.append((c, None))
            continue
        if c.balance is None:  # a balance that isn't a number
            settle(None)
            rows.append(new_row(c, c.direction, f"the balance isn't a readable number ({c.raw_balance})"))
            previous = None
            continue
        settle(c)
    settle(None)
    return rows


# --- The statement ----------------------------------------------------------------------

def parse_spreadsheet(path: Path, *, layouts: dict[str, BankLayout], generic: BankLayout,
                      client_rules: list[ClientRule], settings: Settings, label: str | None = None
                      ) -> list[StatementResult]:
    name = label or path.name
    try:
        sheets = _sheets(path)
    except Exception as exc:
        return [StatementResult(source_file=name, ok=False, transactions=[], status="REVIEW_REQUIRED",
                                error=f"the spreadsheet could not be opened ({exc.__class__.__name__}: {exc}) - "
                                      "save it as .xlsx or .csv and try again")]
    everything = "\n".join(_text_of(cells) for _, cells in sheets)
    tables = [(sheet, cells, t) for sheet, cells in sheets if (t := _find_table(cells))]
    if not tables:
        return [StatementResult(source_file=name, ok=False, transactions=[], status="REVIEW_REQUIRED",
                                error="no table of transactions was found: no row of headings with a date column "
                                      "and an amount (or debit/credit) column, and a balance or description column")]
    try:
        second = dict(_second_sheets(path))
    except Exception:
        second = {}
    results = []
    for sheet, cells, table in tables:
        sheet_label = name if len(tables) == 1 else f"{name} ({sheet})"
        twin = _find_table(second[sheet]) if sheet in second else None
        if twin is not None and (twin.header_row, twin.columns) != (table.header_row, table.columns):
            twin = None
        results.append(_read_table(table, sheet_label, everything, cells, layouts, generic, client_rules,
                                   twin))
    return results


def _opening_above(cells: list[list], header_row: int) -> float | None:
    """An opening balance printed above the table: in the cell with the words
    ("Opening Balance: R1 454.04") or the next cell along. Each cell is read
    whole, so "2 471,87" is one figure."""
    for row in cells[:header_row]:
        for j, value in enumerate(row):
            text = _cell_text(value) if value not in (None, "") else ""
            m = auto_extract._OPENING_WORDS.search(text)
            if not m:
                continue
            rest = text[m.end():].strip(" :")
            if rest:
                amount, _, _ = _money(rest)
                if amount is not None:
                    return amount
            following = next((v for v in row[j + 1:] if v not in (None, "")), None)
            amount, _, _ = _money(following)
            if amount is not None:
                return amount
    return None


def _read_table(table: _Table, label: str, everything: str, cells: list[list], layouts, generic, client_rules,
                second: _Table | None = None):
    layout = detect_bank(everything, layouts, generic)
    period = extract_statement_period(everything, layout, generic)
    account = _find_account_number(everything, layout, generic)
    period_range = _period_range(period)

    source = [_row(table, row_cells) for _, row_cells in table.rows]
    if second is None:
        problem = "the file couldn't be read a second, independent way, so its figures weren't cross-checked"
        for c in source:
            c.note = "not confirmed by a second, independent reading of the file"
    else:
        problem = None
        other = {n: row_cells for n, row_cells in second.rows}
        for (n, _), c in zip(table.rows, source):
            twin = _row(second, other[n]) if n in other else None
            if twin is None or (twin.date, twin.amount, twin.direction, twin.balance) != \
                    (c.date, c.amount, c.direction, c.balance):
                c.note = "the second, independent reading of the file reads this row differently"
    opening = _opening_above(cells, table.header_row)
    sheet_text = _text_of(cells)
    rows = _chain(source, opening)
    reading = assess(rows, opening, sheet_text, layout, generic, groups_confirmed=True, period=period_range)
    if problem:
        reading.problems.append(problem)
    return build_result(reading if reading.rows else None, filename=label, full_text=everything,
                        first_page_text=everything, layout=layout, generic=generic, account_number=account,
                        period=period, client_rules=client_rules, sender=None, subject=None, interactive=False,
                        used_ocr=False, error_hint="no transactions could be read from the table")

